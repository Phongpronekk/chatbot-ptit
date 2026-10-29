"""Build the assessment Word report and deterministic architecture figures."""
from pathlib import Path
import math
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'exports'
ASSETS = OUT / 'architecture_updated'
ASSETS.mkdir(parents=True, exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
INK = '#172B43'


def figure(name, title):
    im = Image.new('RGB', (1800, 1740), 'white')
    d = ImageDraw.Draw(im)
    d.text((65, 34), title, font=ImageFont.truetype(BOLD, 40), fill=INK)
    return im, d


def wrap(draw, text, font, width):
    result = []
    for line in text.split('\n'):
        current = ''
        for word in line.split():
            trial = (current + ' ' + word).strip()
            if draw.textlength(trial, font=font) > width and current:
                result.append(current)
                current = word
            else:
                current = trial
        result.append(current)
    return result


def box(d, xy, title, body, fill='#EEF4FA', border='#7895AF', size=30):
    x1,y1,x2,y2 = xy
    d.rounded_rectangle(xy, radius=18, fill=fill, outline=border, width=3)
    y = y1 + 21
    for text, path, fs in [(title, BOLD, size+3), (body, FONT, size)]:
        font = ImageFont.truetype(path, fs)
        for line in wrap(d,text,font,x2-x1-42):
            d.text((x1+21,y),line,font=font,fill=INK)
            y += fs+11
        y += 10
    assert y < y2+8, (title, y, y2)


def arrow(d, points, label=None, label_pos=None, color='#365D7C'):
    d.line(points, fill=color, width=5, joint='curve')
    a,b=points[-2],points[-1]
    angle=math.atan2(b[1]-a[1],b[0]-a[0])
    d.polygon([b,(b[0]-20*math.cos(angle-.45),b[1]-20*math.sin(angle-.45)),
                 (b[0]-20*math.cos(angle+.45),b[1]-20*math.sin(angle+.45))],fill=color)
    if label:
        p=label_pos or ((points[0][0]+b[0])/2+10,(points[0][1]+b[1])/2)
        font=ImageFont.truetype(FONT,27)
        bounds=d.textbbox(p,label,font=font)
        d.rectangle((bounds[0]-5,bounds[1]-4,bounds[2]+5,bounds[3]+4),fill='white')
        d.text(p,label,font=font,fill=color)


def make_figures():
    im,d=figure('architecture','Kiến trúc các thành phần đã triển khai')
    box(d,(65,120,1735,315),'Giao diện sinh viên   Vue 3 và TypeScript',
        'Pinia lưu trạng thái · Vue Router · Fetch JSON · Tin nhắn chờ và Thinking...\nCookie phiên HttpOnly do BE quản lý · FE không giữ API key')
    arrow(d,[(900,315),(900,390)],'HTTP /api qua Vite proxy',(930,333))
    box(d,(65,390,1735,620),'FastAPI   Điều phối tại BE/main.py',
        'Kiểm tra input và owner của hội thoại → kiểm tra request_id → gọi Knowledge.answer\nThành công: lưu trọn lượt hỏi đáp · Lỗi dịch vụ: HTTP 503, không lưu lượt dở\nLưu lịch sử đã có; chưa truyền lịch sử vào truy vấn hoặc prompt')
    arrow(d,[(420,620),(420,710)],'Câu hỏi hiện tại',(440,645))
    arrow(d,[(1360,620),(1360,710)],'Đọc và ghi theo owner',(1380,645))
    box(d,(65,710,900,1010),'Truy xuất hybrid local',
        'Chuẩn hóa truy vấn · E5-small 384 chiều\nChroma cosine + BM25 không dấu\nLọc bằng chứng → RRF → Top-K 4\nCache BM25 RAM; vector lỗi dùng BM25',size=29)
    box(d,(1000,710,1735,1010),'SQLite   BE/storage/chat.sqlite3',
        'conversations và messages\nCookie phiên ẩn danh 30 ngày\nNguồn được lưu cùng câu trả lời\nChống ghi trùng khi retry',fill='#EDF7F0',size=29)
    arrow(d,[(420,1010),(420,1110)],'Đoạn nguồn đã lọc',(440,1040))
    box(d,(65,1110,900,1420),'Sinh câu trả lời   Knowledge.answer',
        'Nguồn đánh số + câu hỏi + system prompt\nKhông có nguồn: báo thiếu dữ liệu, bỏ LLM\nCó nguồn: gọi provider trong .env\nKết quả + nguồn → BE → FE',size=29)
    box(d,(1000,1110,1735,1420),'Provider được cấu hình',
        'Gemini 3.5 Flash Lite API qua HTTPS\nGửi câu hỏi và các đoạn tài liệu\nOllama: model local, cổng 11434\nChỉ một nhánh được chọn cho mỗi lượt',fill='#FFF5E8',border='#C49D68',size=29)
    arrow(d,[(900,1275),(1000,1275)],'Gọi',(925,1220))
    box(d,(65,1500,1735,1670),'Lưu trữ tri thức trên máy',
        '160 bài năm 2026 trong 5 file Markdown → 1.459 chunk Chroma\nMỗi bài có URL và ngày đăng riêng. Sau ingest cần khởi động lại BE.',fill='#F4F5F7',size=28)
    im.save(ASSETS/'01_architecture.png')

    im,d=figure('sequence','Luồng xử lý một câu hỏi có RAG')
    steps=[
      ('1   Sinh viên gửi câu hỏi','FE hiển thị ngay tin nhắn, bật Thinking...; gửi message và request_id.'),
      ('2   API kiểm tra phiên và hội thoại','Xác thực owner ẩn danh; request_id đã hoàn tất thì trả lại lượt đã lưu.'),
      ('3   Chuẩn hóa và tìm ứng viên','Unicode, khoảng trắng, alias; E5/Chroma và BM25 không dấu cùng truy xuất.'),
      ('4   Lọc bằng chứng và hợp nhất RRF','Cosine theo ngưỡng 0,80; BM25 có cổng khớp riêng; RRF chọn tối đa 4 nguồn.'),
      ('5   Tạo prompt và gọi LLM','Gemini hoặc Ollama đọc các nguồn đánh số; lịch sử chưa được đưa vào.'),
      ('6   Hoàn tất lượt hỏi đáp','Kiểm tra cấu trúc nguồn → transaction SQLite → trả JSON → FE hiện đáp án.')]
    ys=[120,355,590,825,1060,1295]
    for i,((title,body),y) in enumerate(zip(steps,ys)):
        box(d,(70,y,1160,y+170),title,body,size=29)
        if i<5: arrow(d,[(615,y+170),(615,ys[i+1])])
    box(d,(1250,355,1730,560),'Sai owner hoặc input','404 hoặc 422. Không truy cập lịch sử của phiên khác.',fill='#FFF5E8',size=27)
    arrow(d,[(1160,440),(1250,440)])
    box(d,(1250,825,1730,1040),'Không có nguồn','Trả thiếu dữ liệu, không gọi LLM. Điểm cao chưa chứng minh đủ điều kiện.',fill='#FFF5E8',size=27)
    arrow(d,[(1160,910),(1250,910)])
    box(d,(1250,1110,1730,1335),'Lỗi model hoặc mạng','HTTP 503; không lưu lượt hỏi đáp dở. FE hiện lỗi và cho thử lại.',fill='#FFF5E8',size=27)
    arrow(d,[(1160,1145),(1250,1145)])
    d.text((75,1555),'Giới hạn: chưa reformulation, chưa lọc học kỳ và chưa kiểm tra từng khẳng định sau sinh.',font=ImageFont.truetype(FONT,29),fill=INK)
    d.text((75,1610),'FE đang đợi JSON đầy đủ; chưa streaming token.',font=ImageFont.truetype(FONT,29),fill=INK)
    im.save(ASSETS/'02_question_flow.png')

    im,d=figure('ingest','Luồng cập nhật và phiên bản chỉ mục')
    boxes=[
      ((70,120,1730,320),'1   Crawler bài công khai có ngày đăng năm 2026','Lọc ngày từ API, phân loại, loại trùng → staging → sao lưu → thay đúng 5 file Markdown.'),
      ((70,410,1730,610),'2   Đọc nguồn riêng của từng bài và chia đoạn','Parser source_metadata giữ URL/ngày đăng; heading + recursive splitter cắt 400 ký tự, overlap 60 ký tự.'),
      ((70,700,1730,900),'3   Fingerprint và embedding cache','Hash nội dung, metadata, model và cấu hình. Không đổi thì tái sử dụng; thay đổi thì xây thế hệ mới.'),
      ((70,990,1730,1190),'4   Tạo collection Chroma mới','E5 passage: → vector 384 chiều → cosine collection. Chỉ đổi manifest khi hoàn tất.'),
      ((70,1280,1730,1480),'5   Khởi động lại BE và làm nóng cache BM25','BE nạp collection hiện hành và dựng chỉ mục từ khóa RAM; không cần embedding lại khi chỉ đổi thuật ngữ tìm kiếm.')]
    for i,(xy,title,body) in enumerate(boxes):
        box(d,xy,title,body,size=31)
        if i<4: arrow(d,[(900,xy[3]),(900,boxes[i+1][0][1])])
    d.text((75,1540),'Rebuild lỗi: giữ manifest cũ. Thêm, sửa, xóa và ingest lặp được thử trong kho tạm.',font=ImageFont.truetype(FONT,29),fill=INK)
    d.text((75,1590),'Collection cũ còn trên đĩa; xóa khỏi chỉ mục đang dùng không phải xóa vật lý mọi bản cũ.',font=ImageFont.truetype(FONT,29),fill=INK)
    d.text((75,1640),'Đổi riêng Gemini ↔ Ollama không cần embedding lại dữ liệu.',font=ImageFont.truetype(FONT,29),fill=INK)
    im.save(ASSETS/'03_ingest_flow.png')

    im,d=figure('hybrid','Chi tiết tìm kiếm hybrid và các loại điểm')
    box(d,(70,120,1730,300),'Câu hỏi gốc được giữ để gửi vào prompt',
        'Ví dụ: bao hiem  y te, bảo hiểm y tế hoặc BHYT. Không gọi Gemini ở bước truy xuất.',size=30)
    arrow(d,[(900,300),(900,390)])
    box(d,(70,390,1730,585),'Chuẩn hóa truy vấn   BE/search_text.py',
        'Unicode NFC và khoảng trắng → thuật ngữ tên đầy đủ/viết tắt\nTạo biến thể truy vấn; BM25 chuẩn hóa không dấu cả tài liệu và câu hỏi.',size=30)
    arrow(d,[(450,585),(450,700)])
    arrow(d,[(1350,585),(1350,700)])
    box(d,(70,700,860,1060),'Nhánh ngữ nghĩa   E5 và Chroma',
        'Vector câu hỏi 384 chiều\nCosine từ collection đang hoạt động\nTối đa 16 ứng viên khi Top-K = 4\nNgưỡng cosine 0,80\nMargin bổ sung khi thiếu hỗ trợ từ khóa',size=29)
    box(d,(940,700,1730,1060),'Nhánh từ khóa   BM25 local',
        'Chỉ mục RAM trên cùng các chunk\nKhông dấu; thuật ngữ tương đương\nTối đa 16 ứng viên khi Top-K = 4\nCụm ngắn cần khớp đủ và đúng chuỗi\nCâu dài có cổng bao phủ từ khóa',fill='#EDF7F0',size=29)
    arrow(d,[(450,1060),(450,1140),(900,1140)])
    arrow(d,[(1350,1060),(1350,1140),(900,1140)])
    arrow(d,[(900,1140),(900,1210)])
    box(d,(70,1210,1730,1420),'Lọc bằng chứng và Reciprocal Rank Fusion',
        'RRF cộng 1 / (60 + thứ hạng) từ từng nhánh; không cộng thẳng cosine với BM25.\nKhớp từ khóa mạnh có thể được chọn dù cosine dưới ngưỡng. Không đủ bằng chứng thì trả rỗng.',size=30)
    arrow(d,[(900,1420),(900,1510)])
    box(d,(70,1510,1730,1690),'Đầu ra tối đa 4 nguồn kèm diagnostic',
        'score_type: cosine_similarity hoặc bm25; semantic_score, lexical_score, fusion_score tách riêng.\nRRF chưa phải cross-encoder reranking; điểm không phải xác suất đúng.',size=28)
    im.save(ASSETS/'04_hybrid_search.png')


def rich(paragraph, text):
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        run=paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'): run.bold=True
        elif part.startswith('`'): run.font.name='Consolas';run.font.size=Pt(9)


def table(doc, rows):
    t=doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths={2:[1.55,5.0],3:[1.45,.65,4.45],4:[1.1,1.4,1.4,2.75]}.get(len(rows[0]))
    if len(rows[0])==3 and rows[0][0]=='Chặng':
        widths=[1.15,2.2,3.2]
    if len(rows[0])==3 and rows[0][0]=='File':
        widths=[2.0,3.8,.75]
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,(cell,value) in enumerate(zip(cells,row)):
            if widths: cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243D54' if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF'));tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
            p.paragraph_format.keep_with_next = ri == 0
            rich(p,value)
            for run in p.runs:
                run.font.size=Pt(9.2)
                if ri==0: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        prop=t.rows[ri]._tr.get_or_add_trPr()
        keep=OxmlElement('w:cantSplit');prop.append(keep)
        if ri==0: prop.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def main():
    make_figures()
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.8);section.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    # Explicit empty borders override template residue in Word and renderers.
    title_ppr=doc.styles['Title'].element.get_or_add_pPr()
    title_ppr.append(OxmlElement('w:pBdr'))
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',24),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('PTIT Chatbot   |   ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph('Đánh giá hệ thống chatbot PTIT',style='Title')
    doc.add_paragraph('Đối chiếu yêu cầu nguồn và kiến trúc triển khai',style='Subtitle')
    doc.add_paragraph('Ngày đánh giá 29 tháng 10 năm 2026')
    p=doc.add_paragraph();rich(p,'**Kết luận cập nhật: 7,1/10.** Hệ thống đã có RAG hybrid, 160 bài đăng năm 2026, nguồn riêng theo bài và 36 test BE đạt. Các khoảng thiếu lớn vẫn là ngữ cảnh hội thoại, phản hồi người dùng, metadata hiệu lực và đánh giá chất lượng độc lập. Điểm là đánh giá kỹ thuật đề xuất, không phải tỷ lệ trả lời chính xác hoặc điểm chính thức của học phần.')
    doc.add_heading('Nội dung báo cáo',level=1)
    for text in ['Bốn sơ đồ kiến trúc và luồng xử lý của hệ thống hiện tại','Ma trận điểm cho 12 yêu cầu bắt buộc và 2 khuyến nghị','Dữ liệu 160 bài trong 5 file và 1.459 chunk có nguồn riêng','Kết quả 36 test BE và các phép kiểm tra truy xuất thực','Trả lời câu hỏi còn mở và xác định ưu tiên cải thiện']:
        doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('Phạm vi và bằng chứng',level=1)
    doc.add_paragraph('Đánh giá dựa trên mã nguồn trong D:\\chatbot-ptit và SOURCE_REQUIREMENTS.md do người dùng cung cấp, sau đợt thu thập dữ liệu năm 2026 và tối ưu tìm kiếm. Hash dữ liệu vẫn khớp báo cáo thu thập; các test cập nhật dùng kho tạm. Bản này thay các mô tả hiện trạng đã cũ trong báo cáo trước.')
    doc.add_paragraph('Các hình thể hiện chức năng đang có. Những phần chưa triển khai được ghi rõ trong hình hoặc trong bảng đánh giá, tránh nhầm thiết kế đề xuất với hành vi đã kiểm chứng.')
    for i,(file,title,caption) in enumerate([
        ('01_architecture.png','Kiến trúc tổng thể','Hình 1. Các tầng FE, BE, truy xuất, lưu trữ và provider. Gemini nằm ngoài máy; embedding, Chroma và SQLite ở local.'),
        ('02_question_flow.png','Luồng hỏi đáp và ngoại lệ','Hình 2. Trình tự xử lý cùng các nhánh thiếu nguồn, sai phiên và lỗi model. Hiện chỉ câu hỏi mới nhất được gửi vào RAG.'),
        ('03_ingest_flow.png','Luồng lập chỉ mục và cập nhật','Hình 3. Crawler lọc bài năm 2026, staging và sao lưu trước khi ingest. BE nạp collection mới và làm nóng BM25 khi khởi động lại.'),
        ('04_hybrid_search.png','Tìm kiếm hybrid và hợp nhất thứ hạng','Hình 4. Hai nhánh vector và BM25 tìm ứng viên, lọc bằng chứng rồi hợp nhất RRF. Ngưỡng cosine không áp dụng trực tiếp lên BM25.')]):
        doc.add_page_break();doc.add_heading(title,level=1)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run();run.add_picture(str(ASSETS/file),width=Inches(6.6))
        descr=run._r.xpath('.//wp:docPr')[0];descr.set('descr',caption)
        doc.add_paragraph(caption,style='Caption')
    doc.add_page_break()
    source=(ROOT/'docs/SOURCE_ASSESSMENT_UPDATED.md').read_text(encoding='utf-8').splitlines()
    i=1
    while i<len(source):
        line=source[i].strip()
        if not line: i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(source) and source[i].strip().startswith('|'):
                cells=[c.strip() for c in source[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',c) for c in cells): rows.append(cells)
                i+=1
            table(doc,rows);continue
        if line.startswith('```'):
            i+=1
            while i<len(source) and not source[i].startswith('```'):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
                r=p.add_run(source[i]);r.font.name='Consolas';r.font.size=Pt(9)
                i+=1
            i+=1;continue
        if line.startswith('#'):
            level=min(len(line)-len(line.lstrip('#'))-1,3)
            title=re.sub(r'[^\w\s]',' ',line.lstrip('# ').strip(),flags=re.U)
            title=re.sub(r'\s+',' ',title)
            doc.add_heading(title,level=max(level,1))
        else:
            bullet=line.startswith('- ')
            p=doc.add_paragraph(style='List Bullet' if bullet else 'Normal')
            rich(p,line[2:] if bullet else line)
        i+=1
    doc.core_properties.title='Đánh giá hệ thống chatbot PTIT và kiến trúc triển khai'
    doc.core_properties.subject='Đối chiếu SOURCE_REQUIREMENTS và bằng chứng mã nguồn'
    doc.core_properties.author=''
    target=OUT/'Bao_cao_danh_gia_chatbot_PTIT_cap_nhat.docx'
    doc.save(target)
    print(target)


if __name__=='__main__': main()
