"""Bounded, process-local document memory. Original files are never persisted."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import PurePosixPath
import re
from threading import RLock
from time import monotonic
from uuid import uuid4
from zipfile import ZipFile

if __package__:
    from .search_text import KeywordIndex, fold
else:
    from search_text import KeywordIndex, fold

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_ZIP_BYTES = 40 * 1024 * 1024
MAX_FILES = 5
TTL_SECONDS = 2 * 60 * 60
EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.xls'}


class DocumentError(ValueError):
    def __init__(self, detail: str, status: int = 422):
        super().__init__(detail)
        self.status = status


def clean_filename(name: str) -> str:
    name = PurePosixPath((name or '').replace('\\', '/')).name
    name = ''.join(c for c in name if c.isprintable()).strip()[:180]
    if not name or PurePosixPath(name).suffix.lower() not in EXTENSIONS:
        raise DocumentError('Chỉ hỗ trợ PDF, Word (.docx), Excel (.xlsx, .xls). Với .doc, hãy lưu lại thành .docx.', 415)
    return name


def extract_document(filename: str, data: bytes) -> list[dict]:
    """Extract text with stable page/sheet/row references and explicit size limits."""
    filename = clean_filename(filename)
    if not data:
        raise DocumentError('File rỗng. Hãy chọn file có nội dung.')
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError('Mỗi file tối đa 10 MB.', 413)
    extension = PurePosixPath(filename).suffix.lower()
    parts: list[tuple[str, str]] = []
    text_chars = 0

    def add(location, text):
        nonlocal text_chars
        text = text.replace('\x00', '').strip()
        if text:
            text_chars += len(text)
            if text_chars > MAX_TEXT_CHARS:
                raise DocumentError('Nội dung file vượt 500.000 ký tự. Hãy chia thành các file nhỏ hơn.', 413)
            parts.append((location, text))

    try:
        if extension in {'.docx', '.xlsx'}:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 3000 or sum(item.file_size for item in entries) > MAX_ZIP_BYTES:
                    raise DocumentError('File Office quá lớn sau giải nén. Hãy chia nhỏ file.', 413)
        if extension == '.pdf':
            from pypdf import PdfReader
            if not data.startswith(b'%PDF-'):
                raise DocumentError('Nội dung file không phải PDF hợp lệ.')
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError('PDF có mật khẩu. Hãy bỏ mật khẩu rồi tải lại.')
            if len(reader.pages) > 300:
                raise DocumentError('PDF tối đa 300 trang. Hãy chia nhỏ file.', 413)
            for number, page in enumerate(reader.pages, 1):
                add(f'Trang {number}', page.extract_text() or '')
        elif extension == '.docx':
            from docx import Document
            from docx.table import Table
            document = Document(BytesIO(data))
            for number, block in enumerate(document.iter_inner_content(), 1):
                if isinstance(block, Table):
                    header = ' | '.join(cell.text for cell in block.rows[0].cells) if block.rows else ''
                    for row_number, row in enumerate(block.rows, 1):
                        add(f'Bảng {number}, dòng {row_number}',
                            (f'Tiêu đề cột: {header[:1500]}\n' if row_number > 1 else '') +
                            f'Dòng {row_number}: ' + ' | '.join(cell.text for cell in row.cells))
                else:
                    add(f'Đoạn {number}', block.text)
        elif extension == '.xlsx':
            from openpyxl import load_workbook
            # Preserve formulas as text; never evaluate formulas or external links.
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > 50:
                    raise DocumentError('Excel tối đa 50 sheet.', 413)
                for sheet in workbook.worksheets:
                    if (sheet.max_column or 0) > 200 or (sheet.max_row or 0) > 20_000:
                        raise DocumentError('Mỗi sheet tối đa 20.000 dòng và 200 cột.', 413)
                    # Do not trust producer-supplied dimensions to truncate data.
                    sheet.reset_dimensions()
                    header = ''
                    for number, row in enumerate(sheet.iter_rows(), 1):
                        if number > 20_000 or len(row) > 200:
                            raise DocumentError('Mỗi sheet tối đa 20.000 dòng và 200 cột.', 413)
                        values = [f'{cell.column_letter}={cell.value}' for cell in row if cell.value is not None]
                        if not values:
                            continue
                        text = ' | '.join(values)
                        if not header:
                            header = text[:1500]
                        add(f'Sheet {sheet.title}, dòng {number}',
                            (f'Tiêu đề cột: {header}\n' if number > 1 else '') + f'Dòng {number}: ' + text)
            finally:
                workbook.close()
        else:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if workbook.nsheets > 50:
                    raise DocumentError('Excel tối đa 50 sheet.', 413)
                for sheet in workbook.sheets():
                    if sheet.nrows > 20_000 or sheet.ncols > 200:
                        raise DocumentError('Mỗi sheet tối đa 20.000 dòng và 200 cột.', 413)
                    header = ''
                    for number in range(sheet.nrows):
                        values = []
                        for column in range(sheet.ncols):
                            cell = sheet.cell(number, column)
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                            if value != '':
                                values.append(f'Cột {column + 1}={value}')
                        if values:
                            text = ' | '.join(values)
                            if not header:
                                header = text[:1500]
                            add(f'Sheet {sheet.name}, dòng {number + 1}',
                                (f'Tiêu đề cột: {header}\n' if number else '') + f'Dòng {number + 1}: ' + text)
            finally:
                workbook.release_resources()
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError('Không đọc được file. File có thể bị hỏng, có mật khẩu hoặc sai định dạng.') from exc
    if not parts:
        raise DocumentError('Không tìm thấy nội dung chữ. PDF scan/ảnh cần OCR trước khi tải lên.')

    from langchain_text_splitters import RecursiveCharacterTextSplitter
    splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=150)
    # Pack neighboring paragraphs/rows so small tables remain usable as a whole.
    packed = []
    prefix, start, end, buffer = None, '', '', ''
    for location, text in parts:
        match = re.fullmatch(r'(Đoạn |Bảng \d+, dòng |Sheet .+, dòng )(\d+)', location)
        group = match.group(1) if match else None
        number = match.group(2) if match else ''
        if buffer and (group is None or group != prefix or len(buffer) + len(text) + 1 > 1200):
            packed.append((f'{prefix}{start}' + (f'–{end}' if end != start else ''), buffer))
            buffer = ''
        if group is None:
            packed.append((location, text))
        else:
            if not buffer:
                prefix, start = group, number
            end = number
            buffer = buffer + '\n' + text if buffer else text
    if buffer:
        packed.append((f'{prefix}{start}' + (f'–{end}' if end != start else ''), buffer))
    return [{'location': location, 'content': content}
            for location, text in packed for content in splitter.split_text(text)]


@dataclass
class StoredDocument:
    owner: str
    conversation: str
    metadata: dict
    chunks: list[dict]
    deadline: float


class DocumentMemory:
    def __init__(self):
        self._documents: dict[str, StoredDocument] = {}
        self._lock = RLock()

    def cleanup(self):
        with self._lock:
            expired = [identifier for identifier, doc in self._documents.items() if doc.deadline <= monotonic()]
            for identifier in expired:
                del self._documents[identifier]

    def clear(self):
        with self._lock:
            self._documents.clear()

    def list(self, owner: str, conversation: str):
        with self._lock:
            self.cleanup()
            return [dict(doc.metadata) for doc in self._documents.values()
                    if doc.owner == owner and doc.conversation == conversation]

    def add(self, owner: str, conversation: str, filename: str, size: int, chunks: list[dict]):
        with self._lock:
            self.cleanup()
            if len(self.list(owner, conversation)) >= MAX_FILES:
                raise DocumentError('Mỗi cuộc trò chuyện tối đa 5 file. Hãy xóa bớt file.', 409)
            if sum(doc.owner == owner for doc in self._documents.values()) >= 20:
                raise DocumentError('Bộ nhớ tạm của trình duyệt đã đầy. Hãy xóa bớt file.', 409)
            characters = sum(len(chunk['content']) for chunk in chunks)
            if len(self._documents) >= 100 or characters + sum(
                    doc.metadata['characters'] for doc in self._documents.values()) > 10_000_000:
                raise DocumentError('Bộ nhớ tài liệu đang đầy. Vui lòng thử lại sau.', 503)
            timestamp = datetime.now(timezone.utc)
            identifier = str(uuid4())
            metadata = {'id': identifier, 'filename': clean_filename(filename), 'size': size,
                        'characters': characters, 'chunk_count': len(chunks),
                        'created_at': timestamp.isoformat(),
                        'expires_at': (timestamp + timedelta(seconds=TTL_SECONDS)).isoformat()}
            self._documents[identifier] = StoredDocument(owner, conversation, metadata, chunks, monotonic() + TTL_SECONDS)
            return dict(metadata)

    def delete(self, owner: str, conversation: str, identifier: str | None = None):
        with self._lock:
            self.cleanup()
            identifiers = [key for key, doc in self._documents.items()
                           if doc.owner == owner and doc.conversation == conversation
                           and (identifier is None or key == identifier)]
            for key in identifiers:
                del self._documents[key]
            return bool(identifiers)

    def retrieve(self, owner: str, conversation: str, identifiers: list[str], question: str, history=None):
        with self._lock:
            self.cleanup()
            documents = [self._documents.get(identifier) for identifier in identifiers]
            if any(doc is None or doc.owner != owner or doc.conversation != conversation for doc in documents):
                raise DocumentError('File đã hết hạn, bị xóa hoặc không thuộc cuộc trò chuyện này. Hãy tải lại file.', 410)
            records = []
            for doc in documents:
                for chunk in doc.chunks:
                    source = {'document_id': doc.metadata['id'], 'source_file': doc.metadata['filename'],
                              'source_url': '', 'fetched_date': doc.metadata['created_at'][:10],
                              'topic_group': 'File tải lên', 'location': chunk['location'],
                              'title': f"{doc.metadata['filename']} · {chunk['location']}",
                              'content': chunk['content'],
                              'note': 'Nội dung người dùng tải lên; không phải nguồn PTIT đã xác minh.'}
                    records.append({'text': source['title'] + '\n' + source['content'], 'metadata': source})
        index = KeywordIndex(records)
        ranked, _ = index.search(question, 8)
        # Follow-up questions such as "còn điều kiện đó?" can reuse the last question.
        if not ranked and history:
            previous = next((m['text'] for m in reversed(history) if m['role'] == 'user'), '')
            ranked, _ = index.search(previous + ' ' + question, 8)
        summary = any(term in fold(question) for term in ('tom tat', 'tong hop', 'tong quan', 'noi dung chinh'))
        if summary or not ranked:
            # Spread excerpts across files/sections instead of taking only the first page.
            selected = []
            for doc in documents:
                positions = [i for i, record in enumerate(records)
                             if record['metadata']['document_id'] == doc.metadata['id']]
                count = min(len(positions), max(1, 8 // len(documents)))
                selected.extend(positions[round(i * (len(positions) - 1) / max(count - 1, 1))]
                                for i in range(count))
            ranked = [{'index': i, 'score': 0.0} for i in selected][:8]
        return [{**records[row['index']]['metadata'], 'score': round(row['score'], 5),
                 'score_type': 'bm25'} for row in ranked]
