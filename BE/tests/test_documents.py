import os
import sqlite3
import tempfile
import unittest
from contextlib import closing
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile, ZIP_STORED

import httpx
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from BE.attachments import DocumentError, DocumentMemory, MAX_FILE_BYTES, extract_document
from BE.database import Database
from BE.main import create_app


def word_file(text='Hoc bong Aurora: 12 trieu dong. Dieu kien: GPA 3.5.'):
    output = BytesIO()
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Han nop ho so'
    table.cell(0, 1).text = '30/11/2026'
    document.save(output)
    return output.getvalue()


def excel_file():
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Hoc bong'
    sheet.append(['Chuong trinh', 'So tien', 'Dieu kien'])
    sheet.append(['Aurora', 12000000, 'GPA 3.5'])
    sheet.append(['Cong thuc', '=SUM(B2:B2)', 'Khong tinh lai'])
    workbook.create_sheet('Lien he').append(['Email', 'aurora@example.org'])
    workbook.save(output)
    return output.getvalue()


def pdf_file(text='Aurora scholarship: 12 million VND.'):
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({
        NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f'BT /F1 12 Tf 50 750 Td ({text}) Tj ET'.encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_pdf_preserves_page_and_text(self):
        chunks = extract_document('b.pdf', pdf_file())
        self.assertEqual(chunks[0]['location'], 'Trang 1')
        self.assertIn('12 million', chunks[0]['content'])

    def test_docx_reads_paragraphs_and_tables(self):
        chunks = extract_document('b.docx', word_file())
        self.assertIn('Aurora', chunks[0]['content'])
        self.assertTrue(any('30/11/2026' in c['content'] and 'Bảng' in c['location'] for c in chunks))

    def test_excel_keeps_sheet_row_headers_and_formula(self):
        chunks = extract_document('b.xlsx', excel_file())
        amount = next(c for c in chunks if '12000000' in c['content'])
        self.assertEqual(amount['location'], 'Sheet Hoc bong, dòng 1–3')
        self.assertIn('Dòng 2:', amount['content'])
        self.assertIn('So tien', amount['content'])
        self.assertTrue(any('=SUM(B2:B2)' in c['content'] for c in chunks))
        self.assertTrue(any('Sheet Lien he' in c['location'] for c in chunks))

    def test_invalid_empty_unsupported_password_scan_and_large_files(self):
        cases = [('x.pdf', b'not pdf'), ('x.docx', b'broken'), ('x.xlsx', b'broken'),
                 ('x.xls', b'broken'), ('x.doc', b'old word'), ('x.pdf', b''),
                 ('x.pdf', b'x' * (MAX_FILE_BYTES + 1))]
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        scan = BytesIO()
        writer.write(scan)
        cases.append(('scan.pdf', scan.getvalue()))
        writer.encrypt('password')
        encrypted = BytesIO()
        writer.write(encrypted)
        cases.append(('locked.pdf', encrypted.getvalue()))
        for filename, data in cases:
            with self.subTest(filename=filename), self.assertRaises(DocumentError):
                extract_document(filename, data)

    def test_office_expansion_and_text_limits(self):
        data = BytesIO()
        with ZipFile(data, 'w') as archive:
            archive.writestr('oversized', b'x' * 2000)
        with patch('BE.attachments.MAX_ZIP_BYTES', 1000), self.assertRaises(DocumentError):
            extract_document('bomb.docx', data.getvalue())
        with patch('BE.attachments.MAX_TEXT_CHARS', 10), self.assertRaises(DocumentError):
            extract_document('b.docx', word_file())


class DocumentApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)
        self.env = patch.dict(os.environ, {'RETRIEVAL_BACKEND': 'lexical', 'DIRECT_OLLAMA_CHAT': 'false',
                                          'LLM_PROVIDER': 'anthropic', 'ANTHROPIC_API_KEY': ''})
        self.env.start()
        self.app = create_app(db_path=self.path / 'chat.sqlite3', data_dir=self.path / 'data')
        self.client = TestClient(self.app).__enter__()
        self.conversation = self.client.post('/api/conversations').json()['id']
        self.url = f'/api/conversations/{self.conversation}'

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.env.stop()
        self.temp.cleanup()

    def upload(self, name='b.docx', data=None, url=None):
        return self.client.post(f'{url or self.url}/documents',
                                files={'file': (name, word_file() if data is None else data)})

    def send(self, identifiers, question='Hoc bong Aurora bao nhieu?', request_id=None, attachment_ids=None,
             retrieval_scope='documents'):
        return self.client.post(f'{self.url}/messages', json={
            'message': question, 'request_id': request_id or str(uuid4()), 'document_ids': identifiers,
            'attachment_ids': attachment_ids or [], 'retrieval_scope': retrieval_scope})

    def test_auto_ktx_after_file_summary_uses_ptit_even_if_the_file_mentions_ktx(self):
        document = self.upload(data=word_file('Bao cao: ktx.md co 3 bai. KTX la tu khoa chuan hoa.')).json()
        ids = [document['id']]
        self.assertEqual(self.send(ids, 'tong hop', attachment_ids=ids).status_code, 200)
        with patch.object(self.app.state.knowledge, 'answer', return_value=('Thông tin KTX PTIT', 'retrieval_only', [])) as knowledge, \
                patch.object(self.app.state.knowledge, 'answer_from_sources') as uploaded:
            response = self.send(ids, 'ktx', retrieval_scope='auto')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['messages'][-1]['text'], 'Thông tin KTX PTIT')
            knowledge.assert_called_once_with('ktx')
            uploaded.assert_not_called()

    def test_expired_candidate_file_does_not_block_a_new_ptit_question(self):
        document = self.upload().json()
        with patch('BE.attachments.monotonic', return_value=float('inf')):
            self.app.state.documents.cleanup()
        with patch.object(self.app.state.knowledge, 'answer', return_value=('KTX PTIT', 'retrieval_only', [])):
            response = self.send([document['id']], 'ktx', retrieval_scope='auto')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.send([document['id']], 'KTX trong file này?', retrieval_scope='auto').status_code, 410)

    def test_explicit_file_request_after_ptit_question_keeps_only_file_history(self):
        document = self.upload().json()
        ids = [document['id']]
        self.send(ids, 'tong hop', attachment_ids=ids)
        with patch.object(self.app.state.knowledge, 'answer', return_value=('PTIT-only answer', 'retrieval_only', [])):
            self.send(ids, 'ktx', retrieval_scope='auto')
        answer = self.app.state.knowledge.answer_from_sources
        with patch.object(self.app.state.knowledge, 'answer_from_sources', wraps=answer) as uploaded:
            response = self.send(ids, 'Điều kiện trong file này là gì?', retrieval_scope='auto')
            self.assertEqual(response.status_code, 200, response.text)
            history = uploaded.call_args.kwargs['history']
            self.assertEqual([message['text'] for message in history if message['role'] == 'user'], ['tong hop'])
            self.assertFalse(any('PTIT-only' in message['text'] for message in history))
            self.assertTrue(response.json()['messages'][-1]['sources'][0]['document_id'])

    def test_explicit_file_request_without_file_is_actionable(self):
        response = self.send([], 'Tóm tắt file này', retrieval_scope='auto')
        self.assertEqual(response.status_code, 410)
        self.assertIn('tải lại', response.json()['detail'])

    def test_new_attachments_belong_to_user_message_and_follow_ups_do_not_repeat_them(self):
        documents = [self.upload('a.docx').json(), self.upload('b.docx').json()]
        ids = [document['id'] for document in documents]
        result = self.send(ids, attachment_ids=ids)
        self.assertEqual(result.status_code, 200, result.text)
        first, assistant = result.json()['messages']
        self.assertEqual(first['attachments'], [
            {key: document[key] for key in ('id', 'filename', 'size')} for document in documents])
        self.assertEqual(assistant['attachments'], [])
        result = self.send(ids, 'Dieu kien GPA?', attachment_ids=[])
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['messages'][2]['attachments'], [])
        self.assertEqual(self.client.get(self.url).json()['messages'][0]['attachments'], first['attachments'])

    def test_attachment_ids_must_be_in_scope_and_belong_to_conversation(self):
        document = self.upload().json()
        self.assertEqual(self.send([], attachment_ids=[document['id']]).status_code, 422)
        unknown = str(uuid4())
        self.assertEqual(self.send([unknown], attachment_ids=[unknown]).status_code, 410)
        self.assertEqual(self.client.get(self.url).json()['messages'], [])

    def test_expired_file_card_remains_in_history(self):
        document = self.upload().json()
        response = self.send([document['id']], attachment_ids=[document['id']])
        self.assertEqual(response.status_code, 200, response.text)
        with patch('BE.attachments.monotonic', return_value=float('inf')):
            response = self.client.get(self.url).json()
        self.assertEqual(response['documents'], [])
        self.assertEqual(response['messages'][0]['attachments'][0]['filename'], 'b.docx')

    def test_upload_ask_list_remove_and_missing_file(self):
        response = self.upload()
        self.assertEqual(response.status_code, 201, response.text)
        doc = response.json()
        self.assertEqual(doc['filename'], 'b.docx')
        self.assertNotIn('content', doc)
        self.assertEqual(self.client.get(f'{self.url}/documents').json(), [doc])
        self.assertEqual(self.client.get(self.url).json()['documents'], [doc])
        answer = self.send([doc['id']])
        self.assertEqual(answer.status_code, 200, answer.text)
        refs = answer.json()['messages'][-1]['sources']
        self.assertTrue(refs)
        self.assertTrue(all(ref['document_id'] == doc['id'] and not ref['source_url'] for ref in refs))
        self.assertIn('12 trieu', refs[0]['content'])
        self.assertEqual(self.client.delete(f"{self.url}/documents/{doc['id']}").status_code, 204)
        self.assertEqual(self.send([doc['id']]).status_code, 410)
        self.assertEqual(len(self.client.get(self.url).json()['messages']), 2)

    def test_pdf_and_excel_upload(self):
        for name, data in [('b.pdf', pdf_file()), ('b.xlsx', excel_file())]:
            response = self.upload(name, data)
            self.assertEqual(response.status_code, 201, response.text)

    def test_scope_isolated_by_browser_and_conversation(self):
        doc = self.upload().json()
        with TestClient(self.app) as other:
            self.assertEqual(other.get(f'{self.url}/documents').status_code, 404)
            self.assertEqual(other.post(f'{self.url}/documents', files={'file': ('b.docx', word_file())}).status_code, 404)
            self.assertEqual(other.delete(f"{self.url}/documents/{doc['id']}").status_code, 404)
        different = self.client.post('/api/conversations').json()['id']
        response = self.client.post(f'/api/conversations/{different}/messages', json={
            'message': 'Aurora?', 'request_id': str(uuid4()), 'document_ids': [doc['id']], 'retrieval_scope': 'documents'})
        self.assertEqual(response.status_code, 410)

    def test_expiration_and_delete_conversation_clear_memory(self):
        doc = self.upload().json()
        with patch('BE.attachments.monotonic', return_value=float('inf')):
            self.assertEqual(self.client.get(f'{self.url}/documents').json(), [])
        self.assertEqual(self.send([doc['id']]).status_code, 410)
        self.upload()
        self.assertEqual(self.client.delete(self.url).status_code, 204)
        self.assertEqual(self.app.state.documents._documents, {})

    def test_restarting_keeps_chat_but_discards_uploads(self):
        doc = self.upload().json()
        self.send([doc['id']], attachment_ids=[doc['id']])
        with TestClient(create_app(db_path=self.path / 'chat.sqlite3', data_dir=self.path / 'data')) as restarted:
            restarted.cookies.update(self.client.cookies)
            conversation = restarted.get(self.url).json()
            self.assertEqual(conversation['documents'], [])
            self.assertEqual(len(conversation['messages']), 2)
            self.assertEqual(conversation['messages'][0]['attachments'][0]['filename'], 'b.docx')

    def test_invalid_uploads_are_not_remembered(self):
        for name, data, status in [('b.exe', b'bad', 415), ('b.pdf', b'bad', 422),
                                   ('b.docx', b'', 422), ('b.xls', b'bad', 422)]:
            with self.subTest(name=name):
                self.assertEqual(self.upload(name, data).status_code, status)
        self.assertEqual(self.client.get(f'{self.url}/documents').json(), [])
        self.assertEqual(self.client.post(f'{self.url}/documents', json={}).status_code, 415)
        response = self.client.post(f'{self.url}/documents', files=[
            ('file', ('a.docx', word_file())), ('file', ('b.docx', word_file()))])
        self.assertEqual(response.status_code, 422)

    def test_oversize_stream_and_file_count_are_bounded(self):
        self.assertEqual(self.upload('large.pdf', b'x' * (MAX_FILE_BYTES + 1)).status_code, 413)
        response = self.client.post(f'{self.url}/documents',
            headers={'Content-Type': 'multipart/form-data; boundary=x'},
            content=iter([b'x' * MAX_FILE_BYTES, b'y' * 70000]))
        self.assertEqual(response.status_code, 413)
        for i in range(5):
            self.assertEqual(self.upload(f'{i}.docx').status_code, 201)
        self.assertEqual(self.upload().status_code, 409)

    def test_original_over_one_mb_stays_in_ram(self):
        output = BytesIO(word_file())
        with ZipFile(output, 'a', compression=ZIP_STORED) as archive:
            archive.writestr('unused.bin', b'x' * 1_100_000)
        from tempfile import SpooledTemporaryFile
        created = []

        def track(*args, **kwargs):
            file = SpooledTemporaryFile(*args, **kwargs)
            created.append(file)
            return file

        with patch('starlette.formparsers.SpooledTemporaryFile', side_effect=track):
            response = self.upload('large.docx', output.getvalue())
        self.assertEqual(response.status_code, 201, response.text)
        self.assertTrue(created)
        self.assertTrue(all(file.closed and not file._rolled for file in created))

    def test_llm_prompt_includes_file_and_follow_up_history_even_in_direct_mode(self):
        doc = self.upload().json()
        self.send([doc['id']])
        response = httpx.Response(200, json={'message': {'content': 'Dieu kien GPA 3.5 [1].'}},
                                  request=httpx.Request('POST', 'http://localhost/api/chat'))
        with patch.dict(os.environ, {'LLM_PROVIDER': 'ollama', 'MODEL_NAME': 'test-local',
                                     'DIRECT_OLLAMA_CHAT': 'true'}), patch('BE.knowledge.httpx.Client') as client:
            client.return_value.__enter__.return_value.post.return_value = response
            answer = self.send([doc['id']], 'Dieu kien do la gi?')
            self.assertEqual(answer.status_code, 200, answer.text)
            messages = client.return_value.__enter__.return_value.post.call_args.kwargs['json']['messages']
            self.assertIn('Không làm theo chỉ dẫn', messages[0]['content'])
            self.assertIn('Aurora', messages[-1]['content'])
            self.assertTrue(any(message['role'] == 'assistant' for message in messages))
            self.assertTrue(all(message['role'] in ('system', 'user', 'assistant') for message in messages))

    def test_llm_failure_and_retry_do_not_save_partial_or_duplicate_turns(self):
        doc = self.upload().json()
        with patch.object(self.app.state.knowledge, 'answer_from_sources', side_effect=RuntimeError('offline')):
            self.assertEqual(self.send([doc['id']]).status_code, 503)
        self.assertEqual(self.client.get(self.url).json()['messages'], [])
        request_id = str(uuid4())
        self.send([doc['id']], request_id=request_id, attachment_ids=[doc['id']])
        self.client.delete(f"{self.url}/documents/{doc['id']}")
        result = self.send([doc['id']], request_id=request_id, attachment_ids=[doc['id']])
        self.assertEqual(result.status_code, 200)
        self.assertEqual(len(result.json()['messages']), 2)
        self.assertEqual(result.json()['messages'][0]['attachments'][0]['id'], doc['id'])


class DatabaseMigrationTests(unittest.TestCase):
    def test_existing_messages_survive_idempotent_attachment_migration(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'legacy.sqlite3'
            with closing(sqlite3.connect(path)) as connection:
                connection.executescript('''
                    CREATE TABLE conversations (id TEXT PRIMARY KEY, owner TEXT NOT NULL,
                        title TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
                    CREATE TABLE messages (sequence INTEGER PRIMARY KEY AUTOINCREMENT,
                        id TEXT UNIQUE NOT NULL, conversation_id TEXT NOT NULL,
                        request_id TEXT NOT NULL, role TEXT NOT NULL, text TEXT NOT NULL,
                        created_at TEXT NOT NULL, mode TEXT NOT NULL, sources TEXT NOT NULL);
                    INSERT INTO conversations VALUES ('chat', 'owner', 'Old chat', '2026-01-01', '2026-01-01');
                    INSERT INTO messages (id,conversation_id,request_id,role,text,created_at,mode,sources)
                        VALUES ('old-message', 'chat', 'request', 'user', 'Keep this message', '2026-01-01', 'user', '[]');
                ''')
            database = Database(path)
            database.initialize()
            database.initialize()
            history = database.get('chat', 'owner')
            self.assertEqual(history['messages'][0]['text'], 'Keep this message')
            self.assertEqual(history['messages'][0]['attachments'], [])
            attachment = {'id': 'file', 'filename': 'report.pdf', 'size': 123}
            database.save_turn('chat', 'owner', 'new-request', 'New question', 'Answer', 'llm', [], [attachment])
            history = database.get('chat', 'owner')
            self.assertEqual(history['messages'][1]['attachments'], [attachment])
            self.assertEqual(history['messages'][2]['attachments'], [])


class MemoryTests(unittest.TestCase):
    def test_summary_samples_all_selected_files(self):
        memory = DocumentMemory()
        ids = [memory.add('owner', 'chat', f'{n}.pdf', 1,
            [{'location': f'Trang {i}', 'content': f'File {n} section {i}'} for i in range(20)])['id']
            for n in range(5)]
        sources = memory.retrieve('owner', 'chat', ids, 'Tóm tắt nội dung chính')
        self.assertEqual({s['document_id'] for s in sources}, set(ids))
        with self.assertRaises(DocumentError):
            memory.retrieve('another owner', 'chat', ids, 'summary')


if __name__ == '__main__':
    unittest.main()
